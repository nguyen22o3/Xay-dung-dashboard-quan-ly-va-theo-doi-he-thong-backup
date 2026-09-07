from pptx import Presentation
from pptx.util import Inches

prs = Presentation("Báo cáo tiến độ đồ án tốt nghiệp.pptx.pptx")
for i in [1, 3, 5, 7]: # Slides 2, 4, 6, 8 (0-indexed)
    print(f"--- Slide {i+1} ---")
    for j, shape in enumerate(prs.slides[i].shapes):
        if shape.has_text_frame:
            text = shape.text.strip()
            if text:
                print(f"  Shape {j}: '{text[:20]}...' (Left: {shape.left/914400:.2f}in, Top: {shape.top/914400:.2f}in, W: {shape.width/914400:.2f}in, H: {shape.height/914400:.2f}in)")
