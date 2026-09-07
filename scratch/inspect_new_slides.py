from pptx import Presentation
prs = Presentation("Báo cáo tiến độ đồ án tốt nghiệp.pptx.pptx")
print(f"Total slides: {len(prs.slides)}")
for i, slide in enumerate(prs.slides):
    print(f"--- Slide {i} ---")
    for j, shape in enumerate(slide.shapes):
        if shape.has_text_frame:
            text = shape.text.strip().replace('\n', ' ')
            if text:
                print(f"  Shape {j}: '{text[:30]}...'")
