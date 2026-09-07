from pptx import Presentation
from pptx.util import Inches, Pt

# Create presentation
prs = Presentation()

# Slide 1: Title
slide = prs.slides.add_slide(prs.slide_layouts[0])
title = slide.shapes.title
subtitle = slide.placeholders[1]
title.text = "BÁO CÁO TIẾN ĐỘ ĐỒ ÁN TỐT NGHIỆP"
subtitle.text = "Đề tài: Xây dựng dashboard giám sát và quản lý hệ thống backup\n\nSinh viên thực hiện:\n2212339 - Võ Tôn Gia Đạt - CTK46B\n2212355 - Hoàng Thanh Phong - CTK46A\n2212427 - Đặng Đăng Nguyên - CTK46A"

# Slide 2: Content
slide = prs.slides.add_slide(prs.slide_layouts[1])
slide.shapes.title.text = "NỘI DUNG"
content = slide.placeholders[1].text_frame
content.text = "I/ Ý tưởng thực hiện đề tài\nII/ Kiến trúc hệ thống\nIII/ Nguyên lý hoạt động\nIV/ Kết quả đạt được & Demo\nV/ Kế hoạch tiếp theo"

# Slide 3: Phần I
slide = prs.slides.add_slide(prs.slide_layouts[1])
slide.shapes.title.text = "I/ Ý tưởng thực hiện đề tài"
tf = slide.placeholders[1].text_frame
tf.text = "Vấn đề hiện tại:"
p = tf.add_paragraph()
p.text = "- Hệ thống dữ liệu lớn, việc sao lưu phân tán gây khó khăn cho quản lý."
p.level = 1
p = tf.add_paragraph()
p.text = "- Không có công cụ cảnh báo tức thời khi tiến trình backup thất bại."
p.level = 1
p = tf.add_paragraph()
p.text = "Giải pháp:"
p.level = 0
p = tf.add_paragraph()
p.text = "- Xây dựng Dashboard tập trung quản lý lịch sử sao lưu nhiều nguồn."
p.level = 1
p = tf.add_paragraph()
p.text = "- Sao lưu đa đích: Server, Google Drive, NAS."
p.level = 1
p = tf.add_paragraph()
p.text = "- Cảnh báo tự động đa kênh (Email, Telegram, Discord)."
p.level = 1

# Slide 4: Phần II
slide = prs.slides.add_slide(prs.slide_layouts[1])
slide.shapes.title.text = "II/ Kiến trúc hệ thống"
tf = slide.placeholders[1].text_frame
tf.text = "Frontend:"
p = tf.add_paragraph()
p.text = "- ReactJS, Vite, Tailwind CSS (Giao diện trực quan, Dark/Light mode)."
p.level = 1
p = tf.add_paragraph()
p.text = "Backend:"
p.level = 0
p = tf.add_paragraph()
p.text = "- Golang, GORM, JWT (Xử lý đồng thời cao, bảo mật)."
p.level = 1
p = tf.add_paragraph()
p.text = "Database:"
p.level = 0
p = tf.add_paragraph()
p.text = "- PostgreSQL (Lưu trữ bản ghi backup, lịch biểu, cấu hình)."
p.level = 1
p = tf.add_paragraph()
p.text = "Tích hợp bên ngoài:"
p.level = 0
p = tf.add_paragraph()
p.text = "- Google Drive API, Telegram API, Discord Webhook, Gmail SMTP."
p.level = 1

# Slide 5: Phần III
slide = prs.slides.add_slide(prs.slide_layouts[1])
slide.shapes.title.text = "III/ Nguyên lý hoạt động"
tf = slide.placeholders[1].text_frame
tf.text = "1. Luồng hoạt động Sao lưu (Backup Flow):"
p = tf.add_paragraph()
p.text = "- Chọn nguồn -> Đóng gói (tar.gz) -> Phân phối tới các đích (Drive, NAS, Server) -> Lưu lịch sử."
p.level = 1
p = tf.add_paragraph()
p.text = "2. Luồng kiểm tra & Cảnh báo (Cron & Alert):"
p.level = 0
p = tf.add_paragraph()
p.text = "- Background Job quét lịch biểu (Cron) liên tục."
p.level = 1
p = tf.add_paragraph()
p.text = "- Đánh dấu lỗi nếu quá thời gian dung sai (Grace period)."
p.level = 1
p = tf.add_paragraph()
p.text = "- Kích hoạt Notification gửi cảnh báo đến quản trị viên qua đa kênh."
p.level = 1

# Slide 6: Phần IV
slide = prs.slides.add_slide(prs.slide_layouts[1])
slide.shapes.title.text = "IV/ Kết quả đạt được"
tf = slide.placeholders[1].text_frame
tf.text = "- Xây dựng hoàn chỉnh giao diện Dashboard quản lý Backup."
p = tf.add_paragraph()
p.text = "- Thực hiện sao lưu thành công lên Server và Google Drive."
p.level = 0
p = tf.add_paragraph()
p.text = "- Hệ thống tự động gửi thông báo qua Telegram và Email khi có sự kiện."
p.level = 0
p = tf.add_paragraph()
p.text = "- Quản lý người dùng, phân quyền đăng nhập Admin."
p.level = 0
p = tf.add_paragraph()
p.text = "- Tích hợp công cụ chọn file native hệ thống."
p.level = 0
p = tf.add_paragraph()
p.text = "(-> Nhóm tiến hành Demo phần mềm tại đây)"
p.level = 0

# Slide 7: Phần V
slide = prs.slides.add_slide(prs.slide_layouts[1])
slide.shapes.title.text = "V/ Kế hoạch tiếp theo"
tf = slide.placeholders[1].text_frame
tf.text = "- Tối ưu hoá việc nén và upload các file dung lượng lớn."
p = tf.add_paragraph()
p.text = "- Đóng gói toàn bộ hệ thống bằng Docker để dễ dàng triển khai."
p.level = 0
p = tf.add_paragraph()
p.text = "- Cải thiện giao diện hiển thị thanh tiến trình (progress bar)."
p.level = 0
p = tf.add_paragraph()
p.text = "- Bổ sung hệ thống khôi phục dữ liệu (Restore) từ file đã backup."
p.level = 0
p = tf.add_paragraph()
p.text = "- Viết Unit Test và Integration Test."
p.level = 0

# Slide 8: End
slide = prs.slides.add_slide(prs.slide_layouts[0])
title = slide.shapes.title
title.text = "Xin cảm ơn!"
subtitle = slide.placeholders[1]
subtitle.text = "Cảm ơn quý thầy cô đã lắng nghe.\nChúng em xin mời quý thầy cô đặt câu hỏi."

prs.save('/home/ddnguyen/backup-dashboard/Bao_cao_tien_do_hoan_thien.pptx')
