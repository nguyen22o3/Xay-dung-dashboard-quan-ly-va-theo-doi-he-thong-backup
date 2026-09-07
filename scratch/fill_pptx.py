from pptx import Presentation
from pptx.util import Inches, Pt
from pptx.enum.text import MSO_ANCHOR, PP_ALIGN

prs = Presentation("Báo cáo tiến độ đồ án tốt nghiệp.pptx.pptx")

def add_content(slide, text_lines, left, top, width, height, font_size=32):
    txBox = slide.shapes.add_textbox(Inches(left), Inches(top), Inches(width), Inches(height))
    tf = txBox.text_frame
    tf.word_wrap = True
    
    for i, line in enumerate(text_lines):
        if i == 0:
            p = tf.paragraphs[0]
        else:
            p = tf.add_paragraph()
        
        p.text = line
        
        # Simple heuristic to indent bullet points
        if line.startswith('•') or line.startswith('-'):
            p.level = 1
        else:
            p.level = 0
            
        p.font.size = Pt(font_size)

# Slide 2 (NỘI DUNG)
slide2 = prs.slides[1]
lines2 = [
    "I. Ý tưởng thực hiện đề tài",
    "II. Kiến trúc hệ thống",
    "III. Nguyên lý hoạt động",
    "IV. Kết quả đạt được & Demo",
    "V. Kế hoạch tiếp theo"
]
add_content(slide2, lines2, left=2.0, top=4.5, width=16.0, height=6.0, font_size=36)

# Slide 4 (Ý tưởng)
slide4 = prs.slides[3]
lines4 = [
    "Vấn đề hiện tại:",
    "• Hệ thống dữ liệu lớn, việc sao lưu phân tán gây khó khăn cho quản lý.",
    "• Không có công cụ cảnh báo tức thời khi tiến trình backup thất bại.",
    "",
    "Giải pháp:",
    "• Xây dựng Dashboard tập trung quản lý lịch sử sao lưu nhiều nguồn.",
    "• Hỗ trợ sao lưu đa đích: Server nội bộ, Google Drive, NAS.",
    "• Cảnh báo tự động đa kênh (Email, Telegram, Discord)."
]
add_content(slide4, lines4, left=1.23, top=1.8, width=17.5, height=8.5, font_size=32)

# Slide 6 (Kiến trúc)
slide6 = prs.slides[5]
lines6 = [
    "Frontend (Giao diện):",
    "• ReactJS, Vite, Tailwind CSS (Giao diện trực quan, hỗ trợ Dark/Light mode).",
    "",
    "Backend (Xử lý lõi):",
    "• Golang, GORM, JWT (Hiệu năng xử lý đồng thời cao, bảo mật an toàn).",
    "",
    "Database (Lưu trữ):",
    "• PostgreSQL (Lưu trữ các bản ghi backup, lịch biểu, cấu hình hệ thống).",
    "",
    "Tích hợp bên ngoài:",
    "• Google Drive API, Telegram Bot API, Discord Webhook, Gmail SMTP."
]
add_content(slide6, lines6, left=1.23, top=1.8, width=17.5, height=8.5, font_size=30)

# Slide 8 (Nguyên lý hoạt động)
slide8 = prs.slides[7]
lines8 = [
    "1. Luồng hoạt động Sao lưu (Backup Flow):",
    "• Quản trị viên chọn nguồn ➔ Hệ thống nén (tar.gz) ➔ Phân phối tới các đích (Drive, NAS, Server) ➔ Lưu vào lịch sử.",
    "",
    "2. Luồng kiểm tra & Cảnh báo (Cron & Alert):",
    "• Background Job liên tục quét lịch biểu (Cron) đang hoạt động.",
    "• Đánh dấu trạng thái lỗi (Failed/Missed) nếu quá thời gian dung sai (Grace period).",
    "• Kích hoạt Notification gửi cảnh báo đến quản trị viên qua đa kênh."
]
add_content(slide8, lines8, left=1.23, top=1.8, width=17.5, height=8.5, font_size=30)

prs.save("Báo cáo tiến độ đồ án tốt nghiệp_Updated.pptx")
