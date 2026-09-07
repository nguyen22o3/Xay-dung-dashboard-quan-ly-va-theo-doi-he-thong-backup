from pptx import Presentation

prs = Presentation("Báo cáo tiến độ đồ án tốt nghiệp.pptx.pptx")
for i, slide in enumerate(prs.slides):
    print(f"--- Slide {i+1} ---")
    for j, shape in enumerate(slide.shapes):
        if not shape.has_text_frame:
            continue
        text = shape.text.strip()
        print(f"  Shape {j}: {text[:50]}...")
