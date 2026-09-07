from pptx import Presentation

prs = Presentation("Báo cáo tiến độ đồ án tốt nghiệp.pptx.pptx")
for i, slide in enumerate(prs.slides):
    print(f"--- Slide {i+1} ---")
    for j, shape in enumerate(slide.placeholders):
        print(f"  Placeholder {shape.placeholder_format.idx}: type={shape.placeholder_format.type} name='{shape.name}'")
