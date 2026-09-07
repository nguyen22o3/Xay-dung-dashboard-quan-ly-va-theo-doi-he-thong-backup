from pptx import Presentation
from pptx.util import Inches, Pt

prs = Presentation("Báo cáo tiến độ đồ án tốt nghiệp.pptx.pptx")

def add_content(slide, text_lines, left, top, width, height, font_size=32):
    txBox = slide.shapes.add_textbox(Inches(left), Inches(top), Inches(width), Inches(height))
    tf = txBox.text_frame
    tf.word_wrap = True
    
    for i, line in enumerate(text_lines):
        if i == 0:
            p = tf.paragraphs[0]
        else:
            p = tf.add_paragraph()
        
        p.text = line
        
        # Simple heuristic to indent bullet points
        if line.startswith('•') or line.startswith('-'):
            p.level = 1
        else:
            p.level = 0
            
        p.font.size = Pt(font_size)

# Update Slide 8 (Transition IV)
# Shape 1: PHẦN I... -> PHẦN IV
# Shape 2: Ý tưởng... -> Kết quả đạt được
slide8 = prs.slides[8]
for shape in slide8.shapes:
    if shape.has_text_frame:
        text = shape.text.strip().replace('\n', ' ')
        if "PHẦN I" in text:
            shape.text = "PHẦN IV"
        elif "Ý tưởng" in text:
            shape.text = "Kết quả đạt được"

# Update Slide 9 (Content IV)
# Shape 0: Nguyên lý... -> Kết quả đạt được & Demo
# Shape 8: Body -> Replace
slide9 = prs.slides[9]
for shape in slide9.shapes:
    if shape.has_text_frame:
        if "Nguyên lý hoạt động" in shape.text:
            shape.text = "Kết quả đạt được & Demo"

# We can just add the new textbox. The old shape 8 might overlap, so let's delete it.
for shape in slide9.shapes:
    if shape.has_text_frame and "1. Luồng hoạt động Sao lưu" in shape.text:
        # delete shape
        sp = shape._element
        sp.getparent().remove(sp)
        break

lines_iv = [
    "• Xây dựng hoàn chỉnh giao diện Dashboard quản lý Backup (hỗ trợ Dark/Light mode).",
    "• Thực hiện sao lưu thành công từ local lên Server nội bộ và Google Drive.",
    "• Hệ thống tự động phát hiện lỗi và gửi thông báo qua Telegram và Email khi có sự kiện.",
    "• Quản lý người dùng, phân quyền đăng nhập Admin (Bảo mật JWT).",
    "• Tích hợp thành công công cụ chọn file (Native Folder Picker) của hệ thống.",
    "",
    "(Nhóm tiến hành Demo phần mềm thực tế tại đây)"
]
add_content(slide9, lines_iv, left=1.23, top=1.8, width=17.5, height=8.5, font_size=30)

# Update Slide 10 (Transition V)
slide10 = prs.slides[10]
for shape in slide10.shapes:
    if shape.has_text_frame:
        text = shape.text.strip().replace('\n', ' ')
        if "PHẦN I" in text:
            shape.text = "PHẦN V"
        elif "Ý tưởng" in text:
            shape.text = "Kế hoạch tiếp theo"

# Update Slide 11 (Content V)
slide11 = prs.slides[11]
for shape in slide11.shapes:
    if shape.has_text_frame:
        if "Nguyên lý hoạt động" in shape.text:
            shape.text = "Kế hoạch tiếp theo"

for shape in slide11.shapes:
    if shape.has_text_frame and "1. Luồng hoạt động Sao lưu" in shape.text:
        sp = shape._element
        sp.getparent().remove(sp)
        break

lines_v = [
    "• Tối ưu hoá việc đóng gói (nén) và upload các file có dung lượng lớn.",
    "• Đóng gói toàn bộ hệ thống bằng Docker để dễ dàng triển khai (Deploy).",
    "• Cải thiện giao diện: Hiển thị thanh tiến trình (progress bar) chi tiết khi đang chạy backup.",
    "• Bổ sung hệ thống khôi phục dữ liệu (Restore) trực tiếp từ file đã backup.",
    "• Viết Unit Test và Integration Test để đảm bảo độ ổn định cho API."
]
add_content(slide11, lines_v, left=1.23, top=1.8, width=17.5, height=8.5, font_size=32)

prs.save("Báo cáo tiến độ đồ án tốt nghiệp.pptx.pptx")
