from pptx import Presentation
prs = Presentation("Báo cáo tiến độ đồ án tốt nghiệp.pptx.pptx")
for i in [6, 7]:
    print(f"Slide {i}")
    for shp in prs.slides[i].shapes:
        print(f" - {shp.shape_type}")
